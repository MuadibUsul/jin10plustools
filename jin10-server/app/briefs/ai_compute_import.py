"""Import the two Silicon Data workbooks into team-scoped observations.

The source workbooks are historical references, not automatically verified live
data. Import preserves missing rows and precision; it never fills blanks with
zero or carries a value across a missing date.
"""
from __future__ import annotations

import argparse
import datetime as dt
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.models import AIComputeObservation, AIComputeRevision, Team, utcnow


def _num(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value))
    except Exception:
        return None
    return result if result.is_finite() and result >= 0 else None


def _day(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def read_token(path: Path) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True)['Sheet1']
    rows = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        day = _day(row[0] if row else None)
        price = _num(row[1] if len(row) > 1 else None)
        if day and price is not None:
            rows.append({'date': day, 'token_price': price})
    return rows


def read_gpu(path: Path) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True)['Sheet2']
    rows = []
    for row in sheet.iter_rows(min_row=4, values_only=True):
        day = _day(row[0] if row else None)
        if not day:
            continue
        values = [_num(row[i] if len(row) > i else None) for i in (1, 2, 3)]
        if all(value is None for value in values):
            continue
        rows.append({'date': day, 'a100': values[0], 'h100': values[1], 'b200': values[2]})
    return rows


def _upsert(db, team_id, kind: str, rows: list[dict]) -> dict[str, int]:
    existing = {row.date: row for row in db.scalars(
        select(AIComputeObservation).where(
            AIComputeObservation.team_id == team_id, AIComputeObservation.kind == kind))}
    inserted = updated = unchanged = 0
    stamp = utcnow()
    def remember(record):
        exists=db.scalar(select(AIComputeRevision.id).where(AIComputeRevision.observation_id==record.id,
            AIComputeRevision.revision==record.revision))
        if exists:return
        db.add(AIComputeRevision(team_id=team_id,observation_id=record.id,date=record.date,
            kind=record.kind,revision=record.revision,source=record.source,
            values={key:str(getattr(record,key)) if getattr(record,key) is not None else None
                    for key in ('token_price','a100','h100','b200')},created_at=record.updated_at))
    for item in rows:
        day = item['date']
        values = {key: value for key, value in item.items() if key != 'date'}
        record = existing.get(day)
        if record is None:
            record=AIComputeObservation(
                team_id=team_id, kind=kind, date=day, token_price=values.get('token_price'),
                a100=values.get('a100'), h100=values.get('h100'), b200=values.get('b200'),
                source='Silicon Data workbook', source_status='imported_workbook',
                revision=1, created_at=stamp, updated_at=stamp)
            db.add(record);db.flush();remember(record);existing[day]=record
            inserted += 1
            continue
        current = {key: getattr(record, key) for key in values}
        remember(record)
        if current == values:
            unchanged += 1
            continue
        for key, value in values.items():
            setattr(record, key, value)
        record.revision += 1
        record.updated_at = stamp
        remember(record)
        updated += 1
    db.flush()
    return {'inserted': inserted, 'updated': updated, 'unchanged': unchanged,
            'rows': len(rows)}


def import_workbooks(db, team_id, token_path: Path, gpu_path: Path) -> dict:
    token_rows = read_token(Path(token_path))
    gpu_rows = read_gpu(Path(gpu_path))
    if not token_rows or not gpu_rows:
        raise ValueError('工作簿中没有可导入的数据')
    return {'token': _upsert(db, team_id, 'token', token_rows),
            'gpu': _upsert(db, team_id, 'gpu', gpu_rows)}


def main() -> None:
    parser = argparse.ArgumentParser(description='Import Token/GPU workbooks for one team')
    parser.add_argument('--team-id', required=True)
    parser.add_argument('--token', required=True)
    parser.add_argument('--gpu', required=True)
    args = parser.parse_args()
    from app.db.session import get_session_factory

    with get_session_factory()() as db:
        team = db.get(Team, args.team_id)
        if team is None:
            raise SystemExit('团队不存在')
        result = import_workbooks(db, team.id, Path(args.token), Path(args.gpu))
        db.commit()
        print(result)


if __name__ == '__main__':
    main()
