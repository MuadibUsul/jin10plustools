"""PLUS 日报 PPT 生成器：DailyIssue → .pptx

流程：加载标准模板 → 分页（按字数自适应）→ 复制原型页并按标准网格排版 → 删除原型页 → 保存。
"""

from __future__ import annotations

import copy
import datetime as dt
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx.slide import Slide

from . import layout
from .layout import emu
from .models import SECTION_LABELS, DailyIssue, DailySection
from .paginate import ImageFrame, Page, Para, TextFrame, paginate
from .textfit import Metrics, wrap
from .pptx_xml import (
    clone_slide,
    delete_slide,
    max_shape_id,
    paragraphs,
    qn,
    remove,
    rename,
    require_shape,
    set_align,
    set_anchor,
    set_font_size,
    set_font_bold,
    set_geometry,
    set_insets,
    set_line_spacing_pts,
    set_paragraph_text,
    set_shape_id,
    set_shape_text,
    set_space_after,
)
from .template import DEFAULT_TEMPLATE, Shapes, load_template

#: 正文框高度余量，防止测算误差导致 PowerPoint 裁切末行
_TEXT_BOX_SLACK = 0.08


@dataclass
class GenerateReport:
    output: Path
    page_count: int
    sections: list[str]
    warnings: list[str]


# ---------- 封面 ----------


def format_cover_date(date: dt.date) -> str:
    return f"{date.year} 年 {date.month} 月 {date.day} 日"


class CoverLayoutError(ValueError):
    pass


def _cover_lines(title: str, width: float) -> list[str]:
    """原版 32pt，长标题优先在标点处断行，保留完整文字。"""
    lines = wrap(title, width, Metrics(32))
    if len(lines) <= 1:
        return lines
    cuts = [i + 1 for i, ch in enumerate(title) if ch in "，、？：" and i >= 4
            and len(wrap(title[:i + 1], width, Metrics(32))) == 1]
    if cuts:
        cut = cuts[-1]
        return [title[:cut], *_cover_lines(title[cut:], width)]
    return lines


def _fill_cover(slide: Slide, date: dt.date, titles: list[str]) -> None:
    date_shape = require_shape(slide, Shapes.COVER_DATE)
    set_shape_text(date_shape, format_cover_date(date))
    # 日期不换行且留足宽度：不同渲染器（PowerPoint / LibreOffice）字宽略有差异
    date_shape.find(".//" + qn("a:bodyPr")).set("wrap", "none")
    set_geometry(date_shape, w=8.0)

    toc = require_shape(slide, Shapes.COVER_TOC)
    proto = copy.deepcopy(toc)
    remove(toc)
    top, bottom, title_x, title_w = 11.372, 23.25, 2.20, 11.37
    line_pt, slack = 48, 0.08
    wrapped = [_cover_lines(title, title_w) for title in titles]
    used = sum(len(lines) * line_pt / 72 for lines in wrapped)
    if used + slack > bottom - top:
        raise CoverLayoutError("首页精选标题过长，原版字号下无法完整放入首页，请精简标题或减少首页精选数量")
    gap = min(14.4 / 72, (bottom - top - used - slack) / max(1, len(titles) - 1))
    next_id = max_shape_id(slide) + 1
    for no, lines in enumerate(wrapped, start=1):
        for is_number in (True, False):
            shape = copy.deepcopy(proto)
            rename(shape, f"cover-{'number' if is_number else 'title'}-{no}")
            set_shape_id(shape, next_id)
            next_id += 1
            ps = paragraphs(shape)
            p = ps[0]
            runs = p.findall(qn("a:r"))
            # 复用模板编号 / 标题各自的字体、字重、字号与颜色。
            run_proto = copy.deepcopy(runs[0] if is_number else runs[-1])
            end_pr = p.find(qn("a:endParaRPr"))
            if end_pr is not None:
                p.remove(end_pr)
            for run in runs:
                p.remove(run)
            for old_p in ps[1:]:
                old_p.getparent().remove(old_p)
            texts = [f"{no:02d}"] if is_number else lines
            for index, text in enumerate(texts):
                if index:
                    br = etree.SubElement(p, qn("a:br"))
                    br.append(copy.deepcopy(run_proto.find(qn("a:rPr"))))
                run = copy.deepcopy(run_proto)
                run.find(qn("a:t")).text = text
                p.append(run)
            if end_pr is not None:
                p.append(end_pr)
            # 第二行与标题文字对齐，编号使用独立文本框。
            set_line_spacing_pts(p, line_pt)
            set_space_after(p, 0)
            set_insets(shape, left=0, right=0, top=0, bottom=0)
            set_anchor(shape, "t")
            set_geometry(shape, x=1.33 if is_number else title_x, y=top,
                         w=0.70 if is_number else title_w,
                         h=len(texts) * line_pt / 72 + slack)
            slide.shapes._spTree.append(shape)
        top += len(lines) * line_pt / 72 + gap


# ---------- 内容页 ----------


def _apply_chrome(slide: Slide) -> None:
    """页眉页脚按标准网格定位（修正手工模板的细微偏差）"""
    L = layout
    set_geometry(require_shape(slide, Shapes.PLUS_TAG), x=L.MARGIN_X, y=L.HEADER_Y, w=L.TAG_W, h=L.HEADER_H)

    running = require_shape(slide, Shapes.RUNNING_TITLE)
    set_geometry(running, x=L.RUNNING_TITLE_X, y=L.HEADER_Y, w=L.RUNNING_TITLE_W, h=L.HEADER_H)
    set_insets(running, left=0, right=0)

    brand = require_shape(slide, Shapes.BRAND_RESEARCH)
    set_geometry(
        brand, x=L.CONTENT_RIGHT - L.BRAND_RESEARCH_W, y=L.HEADER_Y, w=L.BRAND_RESEARCH_W, h=L.HEADER_H
    )
    set_insets(brand, left=0, right=0)
    set_align(brand, "r")

    set_geometry(require_shape(slide, Shapes.HEADER_RULE), x=L.MARGIN_X, y=L.HEADER_RULE_Y, w=L.CONTENT_WIDTH, h=0)

    disclaimer = require_shape(slide, Shapes.FOOTER_DISCLAIMER)
    set_geometry(disclaimer, x=L.MARGIN_X, y=L.FOOTER_DISCLAIMER_Y, w=L.CONTENT_WIDTH, h=L.FOOTER_H)
    set_insets(disclaimer, left=0, right=0)

    url = require_shape(slide, Shapes.FOOTER_URL)
    set_geometry(url, x=L.MARGIN_X, y=L.FOOTER_LINE_Y, w=4.0, h=L.FOOTER_H)
    set_insets(url, left=0, right=0)

    brand_footer = require_shape(slide, Shapes.FOOTER_BRAND)
    set_geometry(
        brand_footer, x=L.CONTENT_RIGHT - L.FOOTER_BRAND_W, y=L.FOOTER_LINE_Y, w=L.FOOTER_BRAND_W, h=L.FOOTER_H
    )
    set_insets(brand_footer, left=0, right=0)
    set_align(brand_footer, "r")


def _fill_section_header(slide: Slide, page: Page) -> None:
    header = require_shape(slide, Shapes.SECTION_HEADER)
    set_geometry(header, x=layout.MARGIN_X, y=layout.SECTION_HEADER_Y)
    no = require_shape(slide, Shapes.SECTION_NO)
    set_insets(no, left=0)
    set_shape_text(no, f"{page.section_no:02d}")
    set_shape_text(require_shape(slide, Shapes.SECTION_TITLE), SECTION_LABELS[page.section])


def _render_text_frame(
    slide: Slide,
    frame: TextFrame,
    proto_shape: etree._Element,
    title_p: etree._Element,
    body_p: etree._Element,
    shape_id: int,
) -> None:
    sp = copy.deepcopy(proto_shape)
    set_shape_id(sp, shape_id)
    rename(sp, f"body-{shape_id}")
    tx_body = sp.find(qn("p:txBody"))
    for p in tx_body.findall(qn("a:p")):
        tx_body.remove(p)

    font = frame.metrics.font_pt
    for para in frame.paras:
        p = copy.deepcopy(title_p if para.style == "title" else body_p)
        set_paragraph_text(p, para.text)
        set_font_size(p, font)
        if para.style == "title":
            set_font_bold(p)
        set_line_spacing_pts(p, frame.metrics.line_height * 72)
        p.find(qn("a:pPr")).set("hangingPunct", "0")  # 与 LibreOffice 换行一致
        set_space_after(p, 12 * frame.metrics.scale)
        tx_body.append(p)

    set_geometry(sp, x=layout.MARGIN_X, y=frame.top, w=layout.CONTENT_WIDTH, h=frame.height + _TEXT_BOX_SLACK)
    set_insets(sp, left=0, right=0, top=0, bottom=0)
    set_anchor(sp, "t")
    slide.shapes._spTree.append(sp)


def _render_page(slide: Slide, page: Page) -> None:
    _apply_chrome(slide)
    if page.is_section_first:
        _fill_section_header(slide, page)

    proto = require_shape(slide, Shapes.BODY)
    ps = paragraphs(proto)
    title_p, body_p = copy.deepcopy(ps[0]), copy.deepcopy(ps[1])
    proto_copy = copy.deepcopy(proto)
    remove(proto)

    next_id = max_shape_id(slide) + 1
    if not page.frames:
        # 该栏目本期没有素材：保留空白正文框（正文样式），编辑可直接点击输入
        empty = TextFrame(top=page.top, metrics=Metrics(layout.FONT_PT), paras=[Para("body", "")])
        _render_text_frame(slide, empty, proto_copy, title_p, body_p, next_id)
        set_geometry(slide.shapes._spTree[-1], h=layout.BODY_BOTTOM - page.top)
    for frame in page.frames:
        if isinstance(frame, TextFrame):
            _render_text_frame(slide, frame, proto_copy, title_p, body_p, next_id)
            next_id += 1
        elif isinstance(frame, ImageFrame):
            pic = slide.shapes.add_picture(
                str(frame.path), emu(frame.left), emu(frame.top), emu(frame.width), emu(frame.height)
            )
            pic.name = f"image-{pic.shape_id}"
            next_id = max(next_id, pic.shape_id + 1)


# ---------- 入口 ----------


def generate(issue: DailyIssue, output: Path, template_path: Path = DEFAULT_TEMPLATE) -> GenerateReport:
    template = load_template(template_path)
    prs = template.prs
    if issue.item_count == 0:
        raise ValueError("日报没有任何素材")
    sections = issue.ordered_sections()

    result = paginate(sections)

    cover = clone_slide(prs, template.cover)
    items = [item for section in sections for item in section.items]
    indices = issue.cover_indices if issue.cover_indices is not None else list(range(min(12, len(items))))
    if len(set(indices)) != len(indices) or any(i < 0 or i >= len(items) for i in indices):
        raise ValueError("首页精选索引无效")
    _fill_cover(cover, issue.date, [items[i].title for i in sorted(indices)])

    for page in result.pages:
        proto = template.section if page.is_section_first else template.cont
        _render_page(clone_slide(prs, proto), page)

    for proto in template.prototypes:
        delete_slide(prs, proto)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return GenerateReport(
        output=output,
        page_count=len(result.pages) + 1,
        sections=[SECTION_LABELS[s.id] for s in sections],
        warnings=result.warnings,
    )
