"""python-pptx 未提供的底层操作：复制幻灯片、删除幻灯片、按名称查找形状、改写文字与位置。"""

from __future__ import annotations

import copy
from collections.abc import Iterator

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.presentation import Presentation
from pptx.slide import Slide

from .layout import emu

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
_R_ATTRS = (f"{{{NS['r']}}}embed", f"{{{NS['r']}}}id", f"{{{NS['r']}}}link")


def qn(tag: str) -> str:
    prefix, local = tag.split(":")
    return f"{{{NS[prefix]}}}{local}"


# ---------- 幻灯片 ----------


def clone_slide(prs: Presentation, source: Slide) -> Slide:
    """在末尾追加 source 的副本（形状、背景、图片与超链接关系）"""
    new = prs.slides.add_slide(source.slide_layout)
    tree = new.shapes._spTree
    for shape in list(new.shapes):
        tree.remove(shape._element)

    rid_map: dict[str, str] = {}
    for rel in source.part.rels.values():
        if rel.reltype in (RT.NOTES_SLIDE, RT.SLIDE_LAYOUT):
            continue
        if rel.is_external:
            rid_map[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)

    src_csld = source._element.find(qn("p:cSld"))
    new_csld = new._element.find(qn("p:cSld"))
    src_bg = src_csld.find(qn("p:bg"))
    if src_bg is not None:
        new_csld.insert(0, copy.deepcopy(src_bg))

    for child in source.shapes._spTree.iterchildren():
        if child.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        clone = copy.deepcopy(child)
        for el in clone.iter():
            for attr in _R_ATTRS:
                rid = el.get(attr)
                if rid in rid_map:
                    el.set(attr, rid_map[rid])
        tree.append(clone)
    return new


def delete_slide(prs: Presentation, slide: Slide) -> None:
    id_list = prs.slides._sldIdLst
    for sld_id in list(id_list):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            id_list.remove(sld_id)
            return


# ---------- 形状 ----------


def iter_named(root: etree._Element) -> Iterator[tuple[str, etree._Element]]:
    """遍历 root 下所有带 cNvPr 名称的形状（含组合内部）"""
    for el in root.iter(qn("p:sp"), qn("p:grpSp"), qn("p:pic"), qn("p:cxnSp")):
        c_nv_pr = el.find(".//" + qn("p:cNvPr"))
        if c_nv_pr is not None:
            yield c_nv_pr.get("name", ""), el


def find_shape(slide: Slide, name: str) -> etree._Element | None:
    for n, el in iter_named(slide.shapes._spTree):
        if n == name:
            return el
    return None


def require_shape(slide: Slide, name: str) -> etree._Element:
    el = find_shape(slide, name)
    if el is None:
        raise ValueError(f"模板缺少形状「{name}」")
    return el


def rename(el: etree._Element, name: str) -> None:
    el.find(".//" + qn("p:cNvPr")).set("name", name)


def remove(el: etree._Element) -> None:
    el.getparent().remove(el)


def _xfrm(el: etree._Element) -> etree._Element:
    """形状自身的 xfrm（组合用 grpSpPr，普通形状用 spPr）"""
    if el.tag == qn("p:grpSp"):
        return el.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    sp_pr = el.find(qn("p:spPr"))
    return sp_pr.find(qn("a:xfrm"))


def set_geometry(
    el: etree._Element,
    x: float | None = None,
    y: float | None = None,
    w: float | None = None,
    h: float | None = None,
) -> None:
    """设置位置与尺寸（英寸）；未指定的保持原值"""
    xfrm = _xfrm(el)
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    if x is not None:
        off.set("x", str(emu(x)))
    if y is not None:
        off.set("y", str(emu(y)))
    if w is not None:
        ext.set("cx", str(emu(w)))
    if h is not None:
        ext.set("cy", str(emu(h)))


def set_insets(
    el: etree._Element,
    left: float | None = None,
    right: float | None = None,
    top: float | None = None,
    bottom: float | None = None,
) -> None:
    """设置文本框内边距（英寸）；未指定的保持原值"""
    body_pr = el.find(".//" + qn("a:bodyPr"))
    if body_pr is None:
        return
    for attr, value in (("lIns", left), ("rIns", right), ("tIns", top), ("bIns", bottom)):
        if value is not None:
            body_pr.set(attr, str(emu(value)))


def set_anchor(el: etree._Element, anchor: str) -> None:
    body_pr = el.find(".//" + qn("a:bodyPr"))
    if body_pr is not None:
        body_pr.set("anchor", anchor)


def max_shape_id(slide: Slide) -> int:
    ids = [int(e.get("id", "0")) for e in slide.shapes._spTree.iter(qn("p:cNvPr"))]
    return max(ids, default=1)


def set_shape_id(el: etree._Element, shape_id: int) -> None:
    el.find(".//" + qn("p:cNvPr")).set("id", str(shape_id))


def set_align(el: etree._Element, algn: str) -> None:
    for p in el.iter(qn("a:p")):
        p_pr = p.find(qn("a:pPr"))
        if p_pr is None:
            p_pr = etree.SubElement(p, qn("a:pPr"))
            p.insert(0, p_pr)
        p_pr.set("algn", algn)


# ---------- 文字 ----------


def paragraphs(el: etree._Element) -> list[etree._Element]:
    return list(el.iter(qn("a:p")))


def set_paragraph_text(p: etree._Element, text: str) -> None:
    """保留第一个文字段的格式，替换为 text，删除其余文字段"""
    runs = p.findall(qn("a:r"))
    if not runs:
        raise ValueError("段落没有可复用格式的文字段")
    first, rest = runs[0], runs[1:]
    first.find(qn("a:t")).text = text
    for r in rest:
        p.remove(r)
    for br in p.findall(qn("a:br")):
        p.remove(br)


def set_shape_text(el: etree._Element, text: str) -> None:
    """整个形状只保留第一段，文字替换为 text"""
    ps = paragraphs(el)
    set_paragraph_text(ps[0], text)
    for p in ps[1:]:
        p.getparent().remove(p)


def set_font_size(p: etree._Element, pt: float) -> None:
    for tag in ("a:rPr", "a:endParaRPr"):
        for r_pr in p.iter(qn(tag)):
            r_pr.set("sz", str(round(pt * 100)))


def set_font_bold(p: etree._Element) -> None:
    """Record title weight explicitly so PPT and PDF converters select the same face."""
    for tag in ("a:rPr", "a:endParaRPr"):
        for r_pr in p.iter(qn(tag)):
            r_pr.set("b", "1")


def set_space_after(p: etree._Element, pt: float) -> None:
    p_pr = p.find(qn("a:pPr"))
    if p_pr is None:
        return
    spc = p_pr.find(qn("a:spcAft"))
    if spc is None:
        return
    pts = spc.find(qn("a:spcPts"))
    if pts is not None:
        pts.set("val", str(round(pt * 100)))


def set_line_spacing_pts(p: etree._Element, pt: float) -> None:
    """固定行距（磅）。比百分比行距更确定：LibreOffice 对百分比行距会在首行上方额外留白"""
    p_pr = p.find(qn("a:pPr"))
    if p_pr is None:
        p_pr = etree.Element(qn("a:pPr"))
        p.insert(0, p_pr)
    old = p_pr.find(qn("a:lnSpc"))
    ln = etree.Element(qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPts")).set("val", str(round(pt * 100)))
    if old is not None:
        old.addprevious(ln)
        p_pr.remove(old)
    else:
        p_pr.insert(0, ln)
