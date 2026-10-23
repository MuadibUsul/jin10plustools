"""Editable geometry for generated daily PPT slides.

Only article, picture and cover-list shapes are editable. The page chrome stays
fixed, and every edit is checked against the source shape before rendering.
"""
from __future__ import annotations

import hashlib
import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Pt

from app.ppt.pptx_xml import qn
from app.schemas import LayoutOverride


def _editable(shape) -> bool:
    return shape.name.startswith(('body-', 'image-', 'cover-title-', 'cover-number-'))


def _fingerprint(shape) -> str:
    payload = shape.name.encode('utf-8')
    if shape.has_text_frame:
        payload += shape.text.encode('utf-8')
    elif shape.name.startswith('image-'):
        payload += shape.image.blob
    return hashlib.sha256(payload).hexdigest()


def _font_size(shape) -> float | None:
    if not shape.has_text_frame:
        return None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                return round(run.font.size.pt, 2)
    return None


def _replace_paragraph(paragraph, text: str) -> None:
    """Keep the paragraph's first run style and represent newlines as soft breaks."""
    xml = paragraph._p
    runs = xml.findall(qn('a:r'))
    if not runs:
        raise ValueError('该文字框没有可编辑的样式')
    prototype = copy.deepcopy(runs[0])
    for node in list(xml):
        if node.tag in (qn('a:r'), qn('a:br')):
            xml.remove(node)
    end = xml.find(qn('a:endParaRPr'))
    position = list(xml).index(end) if end is not None else len(xml)
    for index, line in enumerate(text.replace('\v', '\n').split('\n')):
        if index:
            br = etree.Element(qn('a:br'))
            properties = prototype.find(qn('a:rPr'))
            if properties is not None:
                br.append(copy.deepcopy(properties))
            xml.insert(position, br)
            position += 1
        run = copy.deepcopy(prototype)
        run.find(qn('a:t')).text = line
        xml.insert(position, run)
        position += 1


def apply_layout(pptx: Path, overrides: list[LayoutOverride]) -> None:
    if not overrides:
        return
    presentation = Presentation(str(pptx))
    seen: set[tuple[int, int]] = set()
    for item in overrides:
        key = (item.page, item.shape_id)
        if key in seen:
            raise ValueError('同一排版元素不能重复修改')
        seen.add(key)
        if item.page > len(presentation.slides):
            raise ValueError('排版页码已变化，请重新进入编辑界面')
        slide = presentation.slides[item.page - 1]
        shape = next((s for s in slide.shapes if s.shape_id == item.shape_id), None)
        if shape is None or not _editable(shape) or _fingerprint(shape) != item.fingerprint:
            raise ValueError('稿件版面已变化，请重新进入编辑界面')
        if item.paragraphs is not None:
            if not shape.has_text_frame or len(item.paragraphs) != len(shape.text_frame.paragraphs):
                raise ValueError('文字段落已变化，请重新进入编辑界面')
            for paragraph, text in zip(shape.text_frame.paragraphs, item.paragraphs):
                _replace_paragraph(paragraph, text)
        if item.font_size_pt is not None:
            if not shape.has_text_frame:
                raise ValueError('图片不能调整字号')
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(item.font_size_pt)
        shape.left = round(item.x * presentation.slide_width)
        shape.top = round(item.y * presentation.slide_height)
        shape.width = round(item.width * presentation.slide_width)
        shape.height = round(item.height * presentation.slide_height)
    presentation.save(str(pptx))


def editable_elements(pptx: Path) -> list[list[dict]]:
    presentation = Presentation(str(pptx))
    width, height = presentation.slide_width, presentation.slide_height
    pages: list[list[dict]] = []
    for page, slide in enumerate(presentation.slides, start=1):
        elements = []
        for shape in slide.shapes:
            if not _editable(shape):
                continue
            elements.append({
                'page': page,
                'shape_id': shape.shape_id,
                'fingerprint': _fingerprint(shape),
                'kind': 'text' if shape.has_text_frame else 'image',
                'label': (shape.text.strip().replace('\n', ' ')[:80] if shape.has_text_frame else '配图') or '空白正文框',
                'x': shape.left / width,
                'y': shape.top / height,
                'width': shape.width / width,
                'height': shape.height / height,
                'font_size_pt': _font_size(shape),
                'paragraphs': ([paragraph.text.replace('\v', '\n') for paragraph in shape.text_frame.paragraphs]
                               if shape.has_text_frame else None),
            })
        pages.append(elements)
    return pages
