"""从手工日报 PPT 还原出一份生成器输入（samples/2026-09-18/issue.json），用于对照测试。

    python scripts/build_sample_issue.py templates/plus_daily_source.pptx samples/2026-09-18

识别规则：思源黑体 Bold 段落为条目标题；「【新闻】」开头为新闻；「【解读】」开头为解读首段；其余为补充段落。
模板中 AI热点、精选观点 为占位（xxxxxxxx），AI热点使用 2026-09-21 从金十页面实测采集的一条素材补上。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SLIDE_SECTIONS = {2: "macro", 3: "macro", 4: "macro", 5: "cn_equity", 6: "global_equity", 9: "charts", 10: "charts"}

LIVE_AI_ITEM = {
    "title": "推理算力占比将大幅提升",
    "news": (
        "《2026年中国人工智能计算力发展评估报告》发布。报告显示，随着智能体技术走向成熟，2026—2030年，"
        "中国活跃智能体复合年均增长率将超150%。报告显示，人工智能正完成从内容生成到任务执行的转变，"
        "全球活跃智能体数量将从2026年的7940万个增长至2030年的22.16亿个。2026年，中国活跃智能体已达到495万个，"
        "据预测，2030年将增长到1.97亿个，复合年均增长率达到151.2%。支撑人工智能应用的算力，其投资规模也在加速扩张："
        "2030年，全球人工智能算力市场规模预计将达到1.25万亿美元，其中，中国市场将达到1501亿美元，均为2025年的4倍。"
    ),
    "interpretation": [
        "报告明确指出，AI正完成从“生成内容”到“执行任务”的转变，从聊天走向干活。智能体执行任务需要持续调用算力，"
        "推理需求将超过训练需求，成为算力消耗的主体。这对算力基础设施的低延迟、高并发、低成本提出了新要求。"
    ],
    "images": [],
}


def is_title(paragraph) -> bool:
    for run in paragraph.runs:
        latin = run._r.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
        if latin is not None and "Bold" in latin.get("typeface", ""):
            return True
    return False


def main(source: Path, out_dir: Path) -> None:
    prs = Presentation(str(source))
    img_dir = out_dir / "images"
    img_dir.mkdir(parents=True, exist_ok=True)
    sections: dict[str, list[dict]] = {}

    for index, slide in enumerate(prs.slides, start=1):
        section = SLIDE_SECTIONS.get(index)
        if not section:
            continue
        items = sections.setdefault(section, [])
        page_items: list[dict] = []
        for shape in slide.shapes:
            if shape.name.startswith("文本框") and shape.has_text_frame and shape.top > 1_000_000:
                for p in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in p.runs).strip()
                    if not text or text == "xxxxxxxx":
                        continue
                    if is_title(p):
                        page_items.append({"title": text, "news": None, "interpretation": [], "images": []})
                    elif text.startswith("【新闻】"):
                        page_items[-1]["news"] = text.removeprefix("【新闻】").strip()
                    elif text.startswith("【解读】"):
                        page_items[-1]["interpretation"].append(text.removeprefix("【解读】").strip())
                    else:
                        page_items[-1]["interpretation"].append(text)
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and page_items:
                name = f"slide{index}.{shape.image.ext}"
                (img_dir / name).write_bytes(shape.image.blob)
                page_items[-1]["images"].append({"path": f"images/{name}"})
        items.extend(page_items)

    sections["ai"] = [LIVE_AI_ITEM]
    order = ["macro", "cn_equity", "global_equity", "ai", "research", "charts"]
    issue = {
        "date": "2026-09-18",
        "sections": [{"id": sid, "items": sections[sid]} for sid in order if sid in sections],
    }
    (out_dir / "issue.json").write_text(json.dumps(issue, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"已写入 {out_dir / 'issue.json'}：" + "，".join(f"{s['id']} {len(s['items'])} 条" for s in issue["sections"]))


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
