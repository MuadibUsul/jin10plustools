from __future__ import annotations

import uuid
from pathlib import Path

import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.ppt.pdf import PdfConversionError
from app.ppt.pptx_xml import qn
from tests.conftest import login
from tests.test_api import INTERPRETATION, material, push


@pytest.fixture(autouse=True)
def fake_pdf(monkeypatch: pytest.MonkeyPatch) -> None:
    def convert(pptx: Path, out_dir: Path) -> Path:
        # 按 PPT 页数生成一份真实的空白 PDF（预览需要能渲染）
        pdf = out_dir / f"{pptx.stem}.pdf"
        doc = pymupdf.open()
        for _ in Presentation(str(pptx)).slides:
            doc.new_page(width=1080, height=1920)
        doc.save(str(pdf))
        return pdf

    monkeypatch.setattr("app.services.issue_builder.pptx_to_pdf", convert)


def seed(client: TestClient) -> tuple[dict[str, str], list[str]]:
    h = login(client, "editor@plus.test")
    items = [
        material(collect_time="2026-09-21T01:00:00Z"),  # 北京时间 9:00
        material(
            source_item_id="flash2", fingerprint="2" * 64, title="推理算力占比将大幅提升", content="报告显示……",
            interpretation_status="fetched", interpretation=INTERPRETATION, collect_time="2026-09-21T02:00:00Z",
        ),
        material(source_item_id="flash3", fingerprint="3" * 64, collect_time="2026-09-19T02:00:00Z"),
    ]
    push(client, h, *items)
    return h, [m["id"] for m in items]


def test_pool_filters_by_beijing_day(client: TestClient) -> None:
    h, ids = seed(client)
    today = client.get("/api/pool", params={"day": "2026-09-21"}, headers=h).json()
    assert {m["id"] for m in today} == set(ids[:2])
    three = client.get("/api/pool", params={"day": "2026-09-21", "days": 3}, headers=h).json()
    assert len(three) == 3


def test_editor_category_is_locked(client: TestClient) -> None:
    h, ids = seed(client)
    res = client.patch(f"/api/materials/{ids[0]}", json={"category": "charts"}, headers=h).json()
    assert res["category"] == "charts" and res["category_locked"] is True
    bad = client.patch(f"/api/materials/{ids[0]}", json={"category": "nope"}, headers=h)
    assert bad.status_code == 422


def test_generate_issue_and_download(client: TestClient, tmp_path: Path) -> None:
    h, ids = seed(client)
    body = {
        "issue_date": "2026-09-21",
        "sections": [{"id": "macro", "material_ids": [ids[0]]}, {"id": "ai", "material_ids": [ids[1]]}],
    }
    res = client.post("/api/issues", json=body, headers=h)
    assert res.status_code == 201, res.text
    issue = res.json()
    assert issue["has_pdf"] and issue["page_count"] == 7

    pdf = client.get(f"/api/issues/{issue['id']}/file", headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")
    assert "attachment" in pdf.headers["content-disposition"]

    pptx = client.get(f"/api/issues/{issue['id']}/file", params={"format": "pptx"}, headers=h)
    path = tmp_path / "test-issue.pptx"
    path.write_bytes(pptx.content)
    # 正文栏目头在组合形状内，按 XML 读取包含组合中的文字。
    texts = "\n".join("".join(node.text or "" for node in s._element.iter(qn("a:t")))
                      for s in Presentation(str(path)).slides)
    # 文字原样进入 PPT
    assert "【新闻】报告显示……" in texts and "【解读】要点一" in texts
    assert all(label in texts for label in ("宏观消息", "AI热点", "精选图表"))

    history = client.get("/api/issues", headers=h).json()
    assert [i["id"] for i in history] == [issue["id"]]


def test_generate_validation_and_permissions(client: TestClient) -> None:
    h, ids = seed(client)
    dup = {"issue_date": "2026-09-21", "sections": [{"id": "macro", "material_ids": [ids[0], ids[0]]}]}
    assert client.post("/api/issues", json=dup, headers=h).status_code == 400
    unknown = {"issue_date": "2026-09-21", "sections": [{"id": "macro", "material_ids": [str(uuid.uuid4())]}]}
    assert client.post("/api/issues", json=unknown, headers=h).status_code == 400
    ok = {"issue_date": "2026-09-21", "sections": [{"id": "macro", "material_ids": [ids[0]]}]}
    assert client.post("/api/issues", json=ok, headers=login(client, "viewer@plus.test")).status_code == 403
    issue = client.post("/api/issues", json=ok, headers=h).json()
    assert client.get(f"/api/issues/{issue['id']}/file", headers=login(client, "out@plus.test")).status_code == 404


def test_pdf_failure_still_returns_pptx(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    def fail(_pptx: Path, _out: Path) -> Path:
        raise PdfConversionError("未安装 LibreOffice")

    monkeypatch.setattr("app.services.issue_builder.pptx_to_pdf", fail)
    h, ids = seed(client)
    issue = client.post(
        "/api/issues", json={"issue_date": "2026-09-21", "sections": [{"id": "macro", "material_ids": [ids[0]]}]}, headers=h
    ).json()
    assert issue["has_pdf"] is False and any("PDF 生成失败" in w for w in issue["warnings"])
    assert client.get(f"/api/issues/{issue['id']}/file", headers=h).status_code == 404


def test_web_page_served(client: TestClient) -> None:
    page = client.get("/daily/")
    assert page.status_code == 200 and "PLUS 日报选稿" in page.text
    js = client.get("/app.js")
    assert js.status_code == 200 and js.headers["cache-control"] == "no-cache"
    guide = client.get("/guide.html")
    assert guide.status_code == 200 and "PLUS 日报使用教程" in guide.text and 'href="/guide.html"' in page.text
    assert client.get("/guide/overview.webp").status_code == 200


def test_edits_are_used_in_ppt_and_can_be_cleared(client: TestClient, tmp_path: Path) -> None:
    h, ids = seed(client)
    edits = {"title": "编辑后的标题", "news": "编辑后的新闻", "interpretation": ["第一段", "第二段"]}
    res = client.put(f"/api/materials/{ids[1]}/edits", json=edits, headers={**h,"If-Match":client.get(f"/api/materials/{ids[1]}",headers=h).json()["updated_at"]}).json()
    assert res["edits"] == edits and res["title"] == "推理算力占比将大幅提升"  # 原文不变

    body = {"issue_date": "2026-09-21", "sections": [{"id": "ai", "material_ids": [ids[1]]}]}
    issue = client.post("/api/issues", json=body, headers=h).json()
    pptx = client.get(f"/api/issues/{issue['id']}/file", params={"format": "pptx"}, headers=h)
    path = tmp_path / "test-edits.pptx"
    path.write_bytes(pptx.content)
    texts = "\n".join(sh.text_frame.text for s in Presentation(str(path)).slides for sh in s.shapes if sh.has_text_frame)
    assert "编辑后的标题" in texts and "【新闻】编辑后的新闻" in texts and "【解读】第一段" in texts and "第二段" in texts
    assert "报告显示" not in texts

    cleared = client.delete(f"/api/materials/{ids[1]}/edits", headers={**h,"If-Match":client.get(f"/api/materials/{ids[1]}",headers=h).json()["updated_at"]}).json()
    assert cleared["edits"] is None


def test_edit_images_must_be_uploaded(client: TestClient) -> None:
    h, ids = seed(client)
    res = client.put(f"/api/materials/{ids[0]}/edits", json={"images": ["0" * 32]}, headers={**h,"If-Match":client.get(f"/api/materials/{ids[0]}",headers=h).json()["updated_at"]})
    assert res.status_code == 400


def test_cover_auto_selection_and_manual_override(client, db_factory):
    from app.db.models import Material, User
    from sqlalchemy import select
    from app.schemas import GenerateRequest
    from app.services.issue_builder import IssueBuildError, _daily_issue

    h, ids = seed(client)
    with db_factory() as db:
        user = db.scalars(select(User).where(User.email == 'editor@plus.test')).one()
        for score, mid in enumerate(ids):
            db.get(Material, uuid.UUID(mid)).importance_score = score
        db.flush()
        req = GenerateRequest.model_validate({'issue_date': '2026-10-14',
              'sections': [{'id': 'ai', 'material_ids': [ids[1]]}, {'id': 'macro', 'material_ids': [ids[0]]}],
              'cover_material_ids': [ids[1]]})
        daily = _daily_issue(db, user, req)
        assert daily.cover_indices == [1]
        assert len(daily.ordered_sections()[0].items) == 1  # 未上首页的稿件仍在正文
        req.cover_material_ids = []
        assert _daily_issue(db, user, req).cover_indices == []
        req.cover_material_ids = [uuid.UUID(ids[2])]
        with pytest.raises(IssueBuildError, match='本期已选素材'):
            _daily_issue(db, user, req)
    bad = {'issue_date': '2026-10-14', 'sections': [{'id': 'macro', 'material_ids': [ids[0]]}],
           'cover_material_ids': [ids[0]] * 13}
    assert client.post('/api/issues/preview', headers=h, json=bad).status_code == 422


def test_reset_ai_layout_unlocks_categories(client: TestClient, db_factory) -> None:
    from app.db.models import Material

    h, ids = seed(client)
    with db_factory() as db:
        m = db.get(Material, uuid.UUID(ids[0]))
        m.ai_details = {"section": {"choice": "macro"}}
        m.category = "macro"
        db.commit()
    client.patch(f"/api/materials/{ids[0]}", json={"category": "charts"}, headers=h)
    rows = client.post("/api/pool/reset-ai", params={"day": "2026-09-21"}, headers=h).json()
    m = next(r for r in rows if r["id"] == ids[0])
    assert m["category"] == "macro" and m["category_locked"] is False


def test_preview_matches_generated_pages(client: TestClient) -> None:
    h, ids = seed(client)
    body = {"issue_date": "2026-09-21", "sections": [{"id": "macro", "material_ids": [ids[0]]}]}
    preview = client.post("/api/issues/preview", json=body, headers=h).json()
    issue = client.post("/api/issues", json=body, headers=h).json()
    assert preview["page_count"] == issue["page_count"] == len(preview["pages"])
    assert all(p.startswith("data:image/png;base64,") for p in preview["pages"])
