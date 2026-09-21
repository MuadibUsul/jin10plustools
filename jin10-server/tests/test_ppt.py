from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from app.ppt import layout, textfit
from app.ppt.cli import load_issue
from app.ppt.generator import format_cover_date, generate
from app.ppt.models import DailyIssue, DailyItem, DailySection, ImageRef, SectionId
from app.ppt.paginate import ImageFrame, TextFrame, paginate, paginate_section
from app.ppt.textfit import Metrics

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "samples" / "2026-09-18" / "issue.json"

# PowerPoint 实测行数（模板第 2 页正文，正文宽 12.85"）——防止测算参数回退
TEMPLATE_MEASURED = [
    ("日本央行加息已在市场预期之内", 1),
    ("【新闻】日本央行加息25个基点，将目标利率由1.00%上调至1.25%，为31年来最高水平，符合市场预期。", 2),
    (
        "【解读】日本央行加息已经是市场的主流预期，结果没有意外。时隔3个月再次加息，是本轮紧缩周期里的最短间隔"
        "（此前多为半年一加）。7-2的投票比例也是日本央行经常出现的票型，包括备受瞩目的2024年7月加息。",
        4,
    ),
    ("日本央行估算中性利率区间1.1%-2.5%，本次1.25% 刚好踩进中性利率下沿，货币政策正式从宽松转向中性区间。", 2),
]


class TestTextfit:
    @pytest.mark.parametrize(("text", "lines"), TEMPLATE_MEASURED)
    def test_never_underestimates_powerpoint(self, text: str, lines: int) -> None:
        # 关闭标点悬挂后只可能多估（多留空白），绝不少估（避免溢出）
        assert lines <= textfit.line_count(text, width_in=12.85) <= lines + 1

    def test_closing_punctuation_never_starts_a_line(self) -> None:
        text = "测" * 25 + "，" + "试" * 10
        wrapped = textfit.wrap(text)
        assert all(not line.startswith("，") for line in wrapped)

    def test_latin_words_are_not_split(self) -> None:
        wrapped = textfit.wrap("测" * 24 + "Anthropic" + "试" * 5)
        assert any("Anthropic" in line for line in wrapped)

    def test_smaller_font_fits_more(self) -> None:
        text = "测" * 600
        assert textfit.line_count(text, Metrics(32)) < textfit.line_count(text, Metrics(36))


def make_image(tmp_path: Path, name: str, w: int, h: int) -> ImageRef:
    path = tmp_path / name
    Image.new("RGB", (w, h), "white").save(path)
    return ImageRef(path=path)


def prose(chars: int, sentence: int = 30) -> str:
    """生成约 chars 字、每句 sentence 字的正文"""
    n = max(1, chars // sentence)
    return "".join("测" * (sentence - 1) + "。" for _ in range(n))


def body_item(title: str, chars: int, images: list[ImageRef] | None = None) -> DailyItem:
    return DailyItem(title=title, news=prose(chars), interpretation=[prose(chars)], images=images or [])


class TestPaginate:
    def test_each_section_starts_new_page_with_header(self) -> None:
        sections = [
            DailySection(id=SectionId.MACRO, items=[body_item("甲", 40)]),
            DailySection(id=SectionId.AI, items=[body_item("乙", 40)]),
        ]
        pages = paginate(sections).pages
        assert [(p.section, p.is_section_first, p.section_no) for p in pages] == [
            (SectionId.MACRO, True, 1),
            (SectionId.AI, True, 2),
        ]

    def test_content_stays_inside_body_area(self) -> None:
        section = DailySection(id=SectionId.MACRO, items=[body_item(f"条目{i}", 180) for i in range(6)])
        pages, _ = paginate_section(section, 1)
        assert len(pages) > 1
        for page in pages:
            assert page.frames[0].top >= page.top - 1e-6
            assert page.cursor <= layout.BODY_BOTTOM + 1e-6

    @pytest.mark.parametrize("unsplittable", [False, True])
    def test_title_is_never_orphaned_at_page_bottom(self, unsplittable: bool) -> None:
        # 覆盖各种字数，使标题落在页尾附近；unsplittable=True 时正文没有句号，无法拆分
        for chars in range(60, 400, 7):
            text = "新" * chars if unsplittable else prose(chars)
            items = [DailyItem(title=f"条目{i}", news=text) for i in range(4)]
            pages, _ = paginate_section(DailySection(id=SectionId.MACRO, items=items), 1)
            for page in pages:
                last = page.frames[-1]
                if isinstance(last, TextFrame):
                    assert last.paras[-1].style != "title", f"chars={chars}"

    def test_image_scales_to_content_width(self, tmp_path: Path) -> None:
        img = make_image(tmp_path, "a.png", 1300, 650)
        section = DailySection(id=SectionId.CHARTS, items=[body_item("图", 20, [img])])
        pages, _ = paginate_section(section, 1)
        frame = next(f for f in pages[0].frames if isinstance(f, ImageFrame))
        assert frame.width == pytest.approx(layout.CONTENT_WIDTH)
        assert frame.height == pytest.approx(layout.CONTENT_WIDTH / 2)
        assert frame.left == pytest.approx(layout.MARGIN_X)

    def test_tall_image_moves_to_next_page_and_fits(self, tmp_path: Path) -> None:
        img = make_image(tmp_path, "tall.png", 600, 1400)
        section = DailySection(id=SectionId.MACRO, items=[body_item("长", 300, [img])])
        pages, _ = paginate_section(section, 1)
        frame = next(f for p in pages for f in p.frames if isinstance(f, ImageFrame))
        assert frame.top + frame.height <= layout.BODY_BOTTOM + 1e-6

    def test_adaptive_font_avoids_nearly_empty_last_page(self) -> None:
        # 36pt 下刚好溢出几行到第二页；缩小字号后一页放下
        chars = 0
        for n in range(300, 900, 10):
            probe = DailySection(id=SectionId.MACRO, items=[DailyItem(title="标题", news=prose(n))])
            pages36 = _pages_at(probe, 36)
            if len(pages36) == 2 and pages36[-1].fill_ratio < layout.SPARSE_LAST_PAGE_RATIO and len(_pages_at(probe, 34)) == 1:
                chars = n
                break
        assert chars, "未找到临界字数"
        section = DailySection(id=SectionId.MACRO, items=[DailyItem(title="标题", news=prose(chars))])
        pages, warnings = paginate_section(section, 1)
        assert len(pages) == 1
        assert pages[0].font_pt == 34
        assert any("34pt" in w for w in warnings)



def _pages_at(section: DailySection, font: float):
    from app.ppt.paginate import _run_section, item_blocks

    blocks = [b for item in section.items for b in item_blocks(item, section.id)]
    return _run_section(blocks, section.id, 1, Metrics(font)).pages


class TestGenerator:
    def test_cover_date_format(self) -> None:
        assert format_cover_date(dt.date(2026, 9, 18)) == "2026 年 9 月 18 日"

    def test_generate_sample_issue(self, tmp_path: Path) -> None:
        issue = load_issue(SAMPLE)
        out = tmp_path / "daily.pptx"
        report = generate(issue, out)

        prs = Presentation(str(out))
        assert len(prs.slides) == report.page_count
        texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]

        cover = texts[0]
        assert "2026 年 9 月 18 日" in cover
        # 目录固定 6 个栏目；精选研报没有素材也保留（空白栏目页）
        assert "05 精选研报" in cover and "06 精选图表" in cover
        assert len(prs.slides) == 1 + 9  # 封面 + 8 页素材 + 精选研报空白页
        assert "可编辑文本" not in cover

        body = "\n".join(texts[1:])
        assert "【新闻】日本央行加息25个基点" in body
        assert "推理算力占比将大幅提升" in body
        assert "条目标题（标题样式）" not in body  # 原型页已删除

    def test_single_section_issue(self, tmp_path: Path) -> None:
        issue = DailyIssue(
            date=dt.date(2026, 9, 21),
            sections=[DailySection(id=SectionId.AI, items=[body_item("仅一条", 50)])],
        )
        out = tmp_path / "one.pptx"
        report = generate(issue, out)
        assert report.page_count == 7  # 封面 + 6 个栏目页（其余 5 个为空白页）
        assert report.sections == ["宏观消息", "国内股市动向", "海外股市动向", "AI热点", "精选研报", "精选图表"]

    def test_empty_issue_rejected(self, tmp_path: Path) -> None:
        with pytest.raises(ValueError):
            generate(DailyIssue(date=dt.date(2026, 9, 21), sections=[]), tmp_path / "x.pptx")
